from __future__ import annotations

import base64
import io
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


MONEY = '$#,##0'
NUMBER = '#,##0.000'
PERCENT = '0.00%'


def _number(value):
    try:
        return float(value or 0)
    except (TypeError, ValueError):
        return 0.0


def _colour(value, fallback):
    text = str(value or "").lstrip("#")
    return text.upper() if re.fullmatch(r"[0-9a-fA-F]{6}", text) else fallback


def _styles(company):
    primary = _colour(company.get("primaryColor"), "123B70")
    accent = _colour(company.get("accentColor"), "E8F0F8")
    return primary, accent, Side(style="thin", color="AAB9C9")


def _logo(ws, company, anchor="A1"):
    match = re.fullmatch(r"data:image/(?:png|jpeg|webp);base64,(.+)", str(company.get("logoDataUrl") or ""), re.I | re.S)
    if not match:
        return
    try:
        image = ExcelImage(io.BytesIO(base64.b64decode(match.group(1))))
        image.height = min(image.height, 55)
        image.width = min(image.width, 150)
        ws.add_image(image, anchor)
    except Exception:
        pass


def _safe_sheet_name(value, used):
    base = re.sub(r"[\\/*?:\[\]]", " ", str(value or "APU")).strip()[:31] or "APU"
    candidate, suffix = base, 2
    while candidate.lower() in used:
        tail = f" ({suffix})"
        suffix += 1
        candidate = f"{base[:31-len(tail)]}{tail}"
    used.add(candidate.lower())
    return candidate


def _set_widths(ws, widths):
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _header(ws, title, company, project, columns=5):
    primary, accent, thin = _styles(company)
    ws.sheet_view.showGridLines = False
    _logo(ws, company)
    ws.merge_cells(start_row=1, start_column=2, end_row=1, end_column=columns)
    cell = ws.cell(1, 2, title)
    cell.font = Font(size=18, bold=True, color=primary)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 45
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=columns)
    ws.cell(2, 1, f"{company.get('name') or 'CodeAPU.cl'}  |  Proyecto: {project.get('name') or ''}")
    ws.cell(2, 1).font = Font(bold=True, color=primary)
    ws.cell(2, 1).fill = PatternFill("solid", fgColor=accent)
    ws.cell(2, 1).alignment = Alignment(horizontal="center")
    ws.cell(2, 1).border = Border(bottom=thin)


def export_apu_workbook(payload: dict, output_path: str | Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    company, project = payload.get("company") or {}, payload.get("project") or {}
    partidas = list(payload.get("partidas") or [])
    used = set()

    summary = wb.create_sheet("Desglose APU")
    used.add("desglose apu")
    _header(summary, "DESGLOSE APU", company, project, 5)
    _set_widths(summary, [18, 58, 14, 18, 20])
    headers = ["Código", "Partida", "Unidad", "Costo directo", "Precio unitario final"]
    for col, value in enumerate(headers, 1):
        summary.cell(4, col, value)
    primary, accent, thin = _styles(company)
    for cell in summary[4]:
        cell.fill = PatternFill("solid", fgColor=primary); cell.font = Font(bold=True, color="FFFFFF")
        cell.alignment = Alignment(horizontal="center"); cell.border = Border(left=thin, right=thin, top=thin, bottom=thin)
    gg_rate = _number(payload.get("generalExpenseRate")) / 100
    utility_rate = _number(payload.get("utilityRate")) / 100
    for row_index, partida in enumerate(partidas, 5):
        direct = sum(_number(row.get("total")) for row in partida.get("rows") or [])
        final = direct * (1 + gg_rate + utility_rate)
        values = [partida.get("code"), partida.get("description"), partida.get("unit"), direct, final]
        for col, value in enumerate(values, 1):
            summary.cell(row_index, col, value).border = Border(bottom=thin)
        summary.cell(row_index, 4).number_format = MONEY; summary.cell(row_index, 5).number_format = MONEY

        ws = wb.create_sheet(_safe_sheet_name(partida.get("code"), used))
        _header(ws, "DESGLOSE APU", company, project, 12)
        _set_widths(ws, [16, 30, 42, 14, 18, 16, 10, 13, 14, 16, 18, 14])
        ws.cell(4, 1, "Código:"); ws.cell(4, 2, partida.get("code"))
        ws.cell(4, 3, "Unidad:"); ws.cell(4, 4, partida.get("unit"))
        ws.cell(5, 1, "Partida:"); ws.merge_cells("B5:L5"); ws.cell(5, 2, partida.get("description"))
        row = 7
        groups = (("M", "1) MATERIALES"), ("O", "2) MANO DE OBRA"), ("E", "3) EQUIPOS Y MAQUINARIAS"), ("S", "4) SUBCONTRATOS"))
        totals = {}
        for prefix, label in groups:
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=12)
            ws.cell(row, 1, label); ws.cell(row, 1).fill = PatternFill("solid", fgColor=primary); ws.cell(row, 1).font = Font(bold=True, color="FFFFFF")
            row += 1
            for col, value in enumerate(("Código PRESTO", "Código recurso", "Descripción", "Familia", "Dependencia", "Destino / Subdestino", "Unidad", "Cantidad", "Factor / Rend.", "P. Unitario", "Importe", "Estado"), 1):
                ws.cell(row, col, value).fill = PatternFill("solid", fgColor=accent); ws.cell(row, col).font = Font(bold=True, color=primary)
            row += 1
            group_rows = [item for item in partida.get("rows") or [] if str(item.get("code") or "").upper().startswith(prefix)]
            subtotal = 0.0
            for item in group_rows:
                subtotal += _number(item.get("total"))
                values = (item.get("prestoCode") or item.get("code"), item.get("codeapuCode"), item.get("description"), item.get("resourceFamilyCode"), item.get("dependency"), f"{item.get('destinationCode') or ''} / {item.get('subdestinationCode') or ''}", item.get("unit"), _number(item.get("quantity")), _number(item.get("factor") if item.get("factor") is not None else 1), _number(item.get("unitPrice")), _number(item.get("total")), item.get("classificationStatus"))
                for col, value in enumerate(values, 1): ws.cell(row, col, value)
                ws.cell(row, 8).number_format = NUMBER; ws.cell(row, 9).number_format = NUMBER; ws.cell(row, 10).number_format = MONEY; ws.cell(row, 11).number_format = MONEY
                row += 1
            totals[prefix] = subtotal
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
            ws.cell(row, 1, f"SUBTOTAL {label.split(') ', 1)[-1]}"); ws.cell(row, 1).font = Font(bold=True); ws.cell(row, 1).alignment = Alignment(horizontal="right")
            ws.cell(row, 11, subtotal); ws.cell(row, 11).number_format = MONEY; ws.cell(row, 11).font = Font(bold=True)
            row += 2
        direct = sum(totals.values()); gg = direct * gg_rate; utility = direct * utility_rate
        for label, value in (("COSTO DIRECTO", direct), (f"GASTOS GENERALES [{gg_rate*100:.2f}%]", gg), (f"UTILIDAD [{utility_rate*100:.2f}%]", utility), ("PRECIO UNITARIO FINAL", direct + gg + utility)):
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=10)
            ws.cell(row, 1, label); ws.cell(row, 1).font = Font(bold=True); ws.cell(row, 1).alignment = Alignment(horizontal="right")
            ws.cell(row, 11, value); ws.cell(row, 11).number_format = MONEY; ws.cell(row, 11).font = Font(bold=True, color=primary)
            row += 1
        ws.freeze_panes = "A7"
    if not partidas:
        summary.cell(5, 1, "No existen APU para exportar.")
    wb.save(output_path)


def export_itemized_workbook(payload: dict, output_path: str | Path) -> None:
    wb = Workbook(); ws = wb.active; ws.title = "Itemizado detallado"
    company, project = payload.get("company") or {}, payload.get("project") or {}
    _header(ws, "ITEMIZADO DETALLADO", company, project, 7)
    _set_widths(ws, [17, 58, 13, 15, 18, 19, 12])
    headers = ("Código", "Descripción", "Unidad", "Cantidad", "P. Unitario", "Total", "Nivel")
    primary, accent, thin = _styles(company)
    for col, value in enumerate(headers, 1):
        cell = ws.cell(4, col, value); cell.fill = PatternFill("solid", fgColor=primary); cell.font = Font(bold=True, color="FFFFFF"); cell.alignment = Alignment(horizontal="center")
    for index, item in enumerate(payload.get("items") or [], 5):
        values = (item.get("code"), item.get("description"), item.get("unit"), _number(item.get("quantity")), _number(item.get("unitPrice")), _number(item.get("total")), item.get("level"))
        for col, value in enumerate(values, 1):
            ws.cell(index, col, value); ws.cell(index, col).border = Border(bottom=thin)
        ws.cell(index, 4).number_format = NUMBER; ws.cell(index, 5).number_format = MONEY; ws.cell(index, 6).number_format = MONEY
        if not str(item.get("unit") or "").strip():
            for cell in ws[index]: cell.fill = PatternFill("solid", fgColor=accent); cell.font = Font(bold=True, color=primary)
    ws.freeze_panes = "A5"; ws.auto_filter.ref = f"A4:G{max(4, ws.max_row)}"
    wb.save(output_path)


def export_view_workbook(payload: dict, output_path: str | Path) -> None:
    wb = Workbook(); ws = wb.active
    ws.title = _safe_sheet_name(payload.get("sheetName") or "Reporte", set())
    company, project = payload.get("company") or {}, payload.get("project") or {}
    columns = list(payload.get("columns") or [])
    _header(ws, str(payload.get("title") or "REPORTE").upper(), company, project, max(1, len(columns)))
    primary, accent, thin = _styles(company)
    row = 4
    for summary in payload.get("summary") or []:
        ws.cell(row, 1, summary.get("label")); ws.cell(row, 2, _number(summary.get("value")))
        if summary.get("type") == "money": ws.cell(row, 2).number_format = MONEY
        row += 1
    row += 1
    chart_payload = payload.get("chart") or {}
    if chart_payload.get("type") == "laborHistogram" and chart_payload.get("months"):
        months = chart_payload["months"]
        ws.cell(row, 1, chart_payload.get("title") or "Histograma mensual de dotación")
        ws.cell(row, 1).font = Font(bold=True, color=primary, size=13)
        chart = BarChart()
        chart.type = "col"; chart.grouping = "stacked"; chart.overlap = 100
        company_label = str(chart_payload.get("companyLabel") or "Empresa oferente no definida")
        chart.title = f"Dotación mensual PROPIA ({company_label}) + SUB"; chart.y_axis.title = "Personas"; chart.x_axis.title = "Mes"
        chart.height = 8.5; chart.width = 26
        matrix_row = row + 18
        ws.cell(matrix_row, 1, "Dotación")
        for col, month in enumerate(months, 2): ws.cell(matrix_row, col, f"M{int(_number(month.get('month')))}")
        for offset, (label, key) in enumerate((("PROPIA", "csv"), ("SUB", "sub"), ("TOTAL", "total")), 1):
            ws.cell(matrix_row + offset, 1, label)
            for col, month in enumerate(months, 2):
                ws.cell(matrix_row + offset, col, int(_number(month.get(key)))).number_format = "0"
        data = Reference(ws, min_col=1, max_col=1 + len(months), min_row=matrix_row + 1, max_row=matrix_row + 2)
        categories = Reference(ws, min_col=2, max_col=1 + len(months), min_row=matrix_row)
        chart.add_data(data, titles_from_data=True, from_rows=True); chart.set_categories(categories)
        if len(chart.series) >= 2:
            chart.series[0].graphicalProperties.solidFill = "D97757"
            chart.series[1].graphicalProperties.solidFill = "C69A45"
        ws.add_chart(chart, f"A{row + 1}")
        for cell in ws[matrix_row]:
            cell.fill = PatternFill("solid", fgColor=primary); cell.font = Font(bold=True, color="FFFFFF"); cell.alignment = Alignment(horizontal="center")
        for data_row in range(matrix_row + 1, matrix_row + 4):
            ws.cell(data_row, 1).fill = PatternFill("solid", fgColor=accent); ws.cell(data_row, 1).font = Font(bold=True, color=primary)
            for col in range(1, 2 + len(months)):
                ws.cell(data_row, col).border = Border(bottom=thin); ws.cell(data_row, col).alignment = Alignment(horizontal="center")
        row = matrix_row + 5
    header_row = row
    for col, column in enumerate(columns, 1):
        cell = ws.cell(row, col, column.get("label")); cell.fill = PatternFill("solid", fgColor=primary); cell.font = Font(bold=True, color="FFFFFF"); cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[get_column_letter(col)].width = max(12, min(55, float(column.get("width") or 18)))
    type_palette = {
        "M": ("D9EAF7", "0B4F8A"),
        "O": ("DCE6F2", "17365D"),
        "E": ("F7DED5", "8B452E"),
        "S": ("E5E7EB", "374151"),
    }
    for record in payload.get("rows") or []:
        row += 1
        for col, column in enumerate(columns, 1):
            value = record.get(column.get("key")); kind = column.get("type")
            cell = ws.cell(row, col, None if value in (None, "") else _number(value) if kind in {"money", "number", "percent"} else value)
            cell.border = Border(bottom=thin)
            if kind == "money": cell.number_format = MONEY
            elif kind == "number": cell.number_format = NUMBER
            elif kind == "percent": cell.number_format = PERCENT
        row_kind = record.get("rowKind")
        if row_kind == "group":
            group_fill = str(record.get("rowColor") or primary).lstrip("#")
            for cell in ws[row]:
                cell.fill = PatternFill("solid", fgColor=group_fill); cell.font = Font(bold=True, color="FFFFFF")
            if len(columns) > 1: ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=len(columns))
        elif row_kind == "typeHeader":
            fill, text = type_palette.get(str(record.get("natureCode") or ""), (accent, primary))
            for cell in ws[row]:
                cell.fill = PatternFill("solid", fgColor=fill); cell.font = Font(bold=True, color=text)
            if len(columns) > 1: ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=len(columns))
        elif row_kind in {"typeSubtotal", "subtotal"}:
            fill, text = type_palette.get(str(record.get("natureCode") or ""), ("E9EEF4", primary))
            for cell in ws[row]:
                cell.fill = PatternFill("solid", fgColor=fill); cell.font = Font(bold=True, color=text)
        elif row_kind == "grandTotal":
            for cell in ws[row]:
                cell.fill = PatternFill("solid", fgColor=primary); cell.font = Font(bold=True, color="FFFFFF")
        elif row_kind == "separator":
            ws.row_dimensions[row].height = 7
            for cell in ws[row]: cell.border = Border()
    ws.freeze_panes = f"A{header_row + 1}"; ws.auto_filter.ref = f"A{header_row}:{get_column_letter(max(1, len(columns)))}{max(header_row, row)}"
    wb.save(output_path)

