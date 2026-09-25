from __future__ import annotations

import csv
import hashlib
import io
import math
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Iterable


HEADER_ALIASES = {
    "code": ("item", "codigo", "codigo item", "cod", "n", "nro", "numero", "partida", "item partida"),
    "description": (
        "nombre partida", "descripcion", "designacion", "denominacion", "detalle",
        "glosa", "especificacion", "nombre", "partida", "concepto", "actividad",
    ),
    "unit": ("unidad", "unid", "uni", "und", "un", "ud", "u m", "um"),
    "quantity": ("cantidad", "cant", "canpres", "metrado", "cubicacion"),
    "price": ("precio unitario", "p unitario", "pu", "p u", "valor unitario", "precio", "unit"),
    "total": ("subtotal", "importe", "total", "monto", "parcial"),
}

SUMMARY_STOP_PREFIXES = (
    "cuadro resumen", "resumen presupuesto", "resumen de presupuesto",
    "resumen costos", "resumen de costos", "resumen final",
)

COMMERCIAL_SUMMARY_PREFIXES = (
    "total costo directo", "gastos generales sobre costo directo",
    "gastos generales obras provisorias", "total gastos generales",
    "utilidad", "utilidades", "costo neto", "iva", "i v a",
    "costo total", "total oferta", "valor total presupuesto",
)

NON_CONCEPT_PREFIXES = (
    "subtotal", "sub total", "total parcial",
)


def plain_text(value: object) -> str:
    text = str(value or "").strip().lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", " ", text).strip()


def bounded_edit_distance(left: str, right: str, limit: int) -> int:
    """Return a Levenshtein distance capped above ``limit`` for short headers."""
    if left == right:
        return 0
    if abs(len(left) - len(right)) > limit:
        return limit + 1
    previous = list(range(len(right) + 1))
    for row_index, left_char in enumerate(left, start=1):
        current = [row_index]
        row_minimum = row_index
        for column_index, right_char in enumerate(right, start=1):
            current.append(min(
                current[-1] + 1,
                previous[column_index] + 1,
                previous[column_index - 1] + (left_char != right_char),
            ))
            row_minimum = min(row_minimum, current[-1])
        if row_minimum > limit:
            return limit + 1
        previous = current
    return previous[-1]


def header_alias_matches(label: str, alias: str) -> bool:
    """Match normalized header labels, tolerating small human typing errors."""
    compact_label = label.replace(" ", "")
    compact_alias = alias.replace(" ", "")
    if label == alias or compact_label == compact_alias:
        return True
    if len(alias) >= 4 and alias in label:
        return True
    # Fuzzy matching is deliberately limited to substantial labels. Short
    # abbreviations such as N, PU or UM would otherwise create false matches.
    shortest = min(len(compact_label), len(compact_alias))
    if shortest < 5:
        return False
    limit = 1 if max(len(compact_label), len(compact_alias)) <= 8 else 2
    return bounded_edit_distance(compact_label, compact_alias, limit) <= limit


def safe_bc3_text(value: object) -> str:
    text = str(value or "").replace("|", " ").replace("\\", " ")
    return re.sub(r"\s+", " ", text).strip()


def parse_number(value: object) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, (int, float)):
        number = float(value)
        return number if math.isfinite(number) else 0.0
    text = str(value).strip().replace("$", "").replace("CLP", "").replace(" ", "")
    text = re.sub(r"[^0-9,.-]", "", text)
    if not text:
        return 0.0
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        parts = text.split(",")
        text = "".join(parts) if len(parts[-1]) == 3 and len(parts) > 1 else text.replace(",", ".")
    elif re.fullmatch(r"-?\d{1,3}(\.\d{3})+", text):
        text = text.replace(".", "")
    try:
        number = float(text)
        return number if math.isfinite(number) else 0.0
    except ValueError:
        return 0.0


def format_number(value: float) -> str:
    if abs(value - round(value)) < 1e-9:
        return str(int(round(value)))
    return f"{value:.6f}".rstrip("0").rstrip(".")


def normalize_code(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value or "").strip().replace(",", ".").replace("/", ".").rstrip("#")
    # En itemizados técnicos es habitual separar la sigla y la numeración con
    # espacios: "AC1  6.2.4" significa AC1.6.2.4, no AC16.2.4.
    text = re.sub(r"(?<=[A-Za-z0-9])\s+(?=[A-Za-z0-9])", ".", text)
    text = re.sub(r"\s+", "", text)
    text = re.sub(r"[^A-Za-z0-9._-]", "", text)
    text = re.sub(r"[._-]+", ".", text).strip(".")
    return text.upper()


def header_map(row: list[object]) -> dict[str, int]:
    result: dict[str, int] = {}
    labels = [plain_text(value) for value in row]
    for field, aliases in HEADER_ALIASES.items():
        for index, label in enumerate(labels):
            if not label:
                continue
            if any(header_alias_matches(label, alias) for alias in aliases):
                result[field] = index
                break
    return result


def parent_code(code: str, available: set[str]) -> str | None:
    parts = code.split(".")
    while len(parts) > 1:
        parts.pop()
        candidate = ".".join(parts)
        if candidate in available:
            return candidate
    # Algunos itemizados usan una jerarquía compacta en el primer bloque:
    # B > B1 > B1.1 o AP > AP1 > AP1.1. Solo se acepta un padre que termine
    # en letra y cuyo sufijo sea exclusivamente numérico, evitando confundir
    # códigos hermanos como A1 y A10.
    compact_candidates = [
        candidate
        for candidate in available
        if candidate != code
        and code.startswith(candidate)
        and candidate[-1:].isalpha()
        and code[len(candidate):].isdigit()
    ]
    if compact_candidates:
        return max(compact_candidates, key=len)
    return None


def natural_key(value: str) -> tuple:
    return tuple(int(part) if part.isdigit() else part.lower() for part in re.split(r"(\d+)", value))


PROJECT_TITLE_WORDS = (
    "proyecto", "obra", "edificio", "sede", "contingencia", "sector",
    "torre", "cesfam", "hospital", "escuela", "colegio", "centro", "condominio",
)

CENTRAL_CHAPTER_LABELS = (
    "obra gruesa", "obras previas", "obras preliminares", "terminaciones",
    "especialidades", "cierre", "proyecto y tramitaciones", "proyecto as built",
)


def row_text_values(row: list[object]) -> list[str]:
    return [safe_bc3_text(value) for value in row if safe_bc3_text(value)]


def initial_project_title(rows: list[list[object]], header_index: int) -> str:
    """Obtiene el nombre de obra más cercano al encabezado sin confundir metadatos."""
    for row in reversed(rows[:header_index]):
        values = row_text_values(row)
        if not values:
            continue
        for value in values:
            normalized = plain_text(value)
            if normalized in {"denominacion", "descripcion", "item", "partida"}:
                continue
            if any(label in normalized for label in ("comuna", "fecha", "n rol", "perm edif", "ubicacion", "superficie")):
                continue
            if any(word in normalized for word in PROJECT_TITLE_WORDS):
                return re.sub(r"^itemizado\s*[:\-]?\s*", "", value, flags=re.IGNORECASE).strip() or value
    return ""


def project_boundary_title(row: list[object], mapping: dict[str, int]) -> str:
    """Detecta rótulos de una nueva obra; exige una fila sin datos de concepto."""
    def mapped(field: str) -> object:
        index = mapping.get(field, -1)
        return row[index] if 0 <= index < len(row) else ""

    description = safe_bc3_text(mapped("description"))
    unit = safe_bc3_text(mapped("unit"))
    if description or unit or parse_number(mapped("quantity")) or parse_number(mapped("price")) or parse_number(mapped("total")):
        return ""
    values = row_text_values(row)
    for value in values:
        normalized = plain_text(value)
        if not normalized or normalized.startswith(("sigla ", "costo total", "total ", "subtotal", "cuadro resumen")):
            continue
        if normalized in CENTRAL_CHAPTER_LABELS:
            continue
        strong = (
            "itemizado" in normalized
            or re.match(r"^(obra\s*:|proyecto\s*:|edificio\b|torre\b|sector\b|sede\b)", normalized)
            or re.match(r"^contingencia(?:\s+\d+|\s+sede)\b", normalized)
        )
        if strong:
            return re.sub(r"^itemizado\s*[:\-]?\s*", "", value, flags=re.IGNORECASE).strip() or value
    return ""


def commercial_summary_boundary(row: list[object], mapping: dict[str, int]) -> bool:
    """Detect the commercial block even when merged cells shifted its labels."""
    code_index = mapping.get("code", -1)
    raw_code = row[code_index] if 0 <= code_index < len(row) else ""
    if safe_bc3_text(raw_code):
        return False
    for value in row_text_values(row):
        normalized = re.sub(r"^\d+\s*", "", plain_text(value))
        if any(normalized.startswith(prefix) for prefix in COMMERCIAL_SUMMARY_PREFIXES):
            return True
    return False


@dataclass
class Item:
    code: str
    description: str
    unit: str = ""
    quantity: float = 0.0
    price: float = 0.0
    total: float = 0.0
    synthetic: bool = False
    bc3_code: str = ""
    original_code: str = ""
    source_row: int = 0
    project_index: int = 1
    project_name: str = ""


@dataclass
class ConversionResult:
    source_name: str
    sheet_name: str
    project_name: str
    items: list[Item]
    warnings: list[str] = field(default_factory=list)
    bc3_bytes: bytes = b""
    root_total: float = 0.0
    leaf_count: int = 0
    chapter_count: int = 0
    normalization_report: dict = field(default_factory=dict)
    project_roots: list[dict] = field(default_factory=list)
    import_payload: dict = field(default_factory=dict)

    def preview(self, limit: int = 10000) -> dict:
        if self.import_payload:
            return {
                "sourceName": self.source_name,
                "sheetName": self.sheet_name,
                "projectName": self.project_name,
                "rootTotal": self.root_total,
                "itemCount": len(self.import_payload.get("items", [])),
                "leafCount": self.leaf_count,
                "chapterCount": self.chapter_count,
                "warnings": self.warnings,
                "normalizationReport": self.normalization_report,
                "previewTruncated": False,
                **self.import_payload,
            }
        available = {item.code for item in self.items}
        parent_by_code = {item.code: parent_code(item.code, available) for item in self.items}
        parents = {parent for parent in parent_by_code.values() if parent}
        bc3_by_original = {item.code: item.bc3_code for item in self.items}
        root_by_code: dict[str, str] = {}
        for item in self.items:
            current = item.code
            while parent_by_code.get(current):
                current = parent_by_code[current]
            root_by_code[item.code] = current
        item_by_code = {item.code: item for item in self.items}

        def hierarchy_level(code: str) -> int:
            level = 2
            seen: set[str] = set()
            current = code
            while parent_by_code.get(current) and current not in seen:
                seen.add(current)
                current = parent_by_code[current]
                level += 1
            return level

        def root_kind(code: str) -> str:
            description = plain_text(item_by_code[root_by_code[code]].description)
            commercial_labels = (
                "costo total", "gastos generales", "utilidad", "valor neto",
                "iva", "i v a", "valor total presupuesto", "total oferta",
            )
            return "commercial" if any(label in description for label in commercial_labels) else "budget"
        project_roots = self.project_roots or [{"code": "PRJ001", "description": self.project_name, "total": self.root_total}]
        project_by_index = {index + 1: root for index, root in enumerate(project_roots)}
        return {
            "sourceName": self.source_name,
            "sheetName": self.sheet_name,
            "projectName": self.project_name,
            "projectRoot": {**project_roots[0], "level": 1},
            "projectRoots": [{**root, "level": 1} for root in project_roots],
            "projectCount": len(project_roots),
            "rootTotal": self.root_total,
            "itemCount": len(self.items),
            "leafCount": self.leaf_count,
            "chapterCount": self.chapter_count,
            "warnings": self.warnings,
            "normalizationReport": self.normalization_report,
            "items": [
                {
                    "code": item.bc3_code,
                    "originalCode": item.original_code or item.code,
                    "description": item.description,
                    "unit": item.unit,
                    "quantity": item.quantity,
                    "price": item.price,
                    "total": item.total,
                    # ROOT## es el nivel 1 (proyecto). Los conceptos superiores
                    # del Excel son capítulos de nivel 2.
                    "level": hierarchy_level(item.code),
                    "synthetic": item.synthetic,
                    "hasChildren": item.code in parents,
                    "isChapter": not bool(item.unit.strip()),
                    "isPartida": bool(item.unit.strip()),
                    "parentCode": bc3_by_original.get(parent_by_code[item.code], project_by_index[item.project_index]["code"]),
                    "rootCode": bc3_by_original.get(root_by_code[item.code], item.bc3_code),
                    "rootKind": root_kind(item.code),
                    "projectCode": project_by_index[item.project_index]["code"],
                    "projectName": project_by_index[item.project_index]["description"],
                }
                for item in self.items[:limit]
            ],
            "previewTruncated": len(self.items) > limit,
        }


def read_workbook(data: bytes, suffix: str) -> list[tuple[str, list[list[object]]]]:
    suffix = suffix.lower()
    if suffix == ".csv":
        text = data.decode("utf-8-sig", errors="replace")
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t|")
        except csv.Error:
            dialect = csv.excel
            dialect.delimiter = ";"
        return [("CSV", [list(row) for row in csv.reader(io.StringIO(text), dialect)])]
    if suffix == ".xls":
        try:
            import xlrd
        except ImportError as exc:
            raise ValueError("Para leer .xls instala xlrd o guarda el archivo como .xlsx.") from exc
        book = xlrd.open_workbook(file_contents=data)
        return [
            (sheet.name, [sheet.row_values(index) for index in range(sheet.nrows)])
            for sheet in book.sheets()
        ]
    if suffix not in {".xlsx", ".xlsm"}:
        raise ValueError("Formato no soportado. Usa .xlsx, .xlsm, .xls o .csv.")
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ValueError("Falta openpyxl. Ejecuta: pip install -r requirements.txt") from exc
    workbook = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    try:
        return [
            (sheet.title, [list(row) for row in sheet.iter_rows(values_only=True)])
            for sheet in workbook.worksheets
        ]
    finally:
        workbook.close()


def parse_sheet(rows: list[list[object]]) -> tuple[list[Item], int, dict[str, int]]:
    best_index = -1
    best_mapping: dict[str, int] = {}
    best_score = -1
    # Se amplía el rango porque algunos mandantes anteponen portadas extensas.
    for index, row in enumerate(rows[:200]):
        mapping = header_map(row)
        score = len(mapping) + (3 if "description" in mapping else 0) + (2 if "code" in mapping else 0)
        if score > best_score:
            best_index, best_mapping, best_score = index, mapping, score
    if "description" not in best_mapping:
        return [], -1, {}
    mapping = best_mapping
    items: list[Item] = []
    auto_number = 1
    active_project_name = initial_project_title(rows, best_index)
    active_project_index = 1
    active_project_items = 0
    pending_project_name = ""
    code_index = mapping.get("code", -1)
    description_index = mapping.get("description", -1)
    regular_codes: set[str] = set()
    if code_index >= 0:
        for candidate_row in rows[best_index + 1 :]:
            raw = candidate_row[code_index] if code_index < len(candidate_row) else ""
            candidate_description = candidate_row[description_index] if 0 <= description_index < len(candidate_row) else ""
            is_subtotal = plain_text(candidate_description).startswith(NON_CONCEPT_PREFIXES)
            if raw and not is_subtotal and not re.search(r"\bSIGLA\b", str(raw), re.IGNORECASE):
                regular_codes.add(normalize_code(raw))
    for row_number, row in enumerate(rows[best_index + 1 :], start=best_index + 2):
        def cell(field: str) -> object:
            idx = mapping.get(field, -1)
            return row[idx] if idx >= 0 and idx < len(row) else ""

        raw_code = cell("code") if "code" in mapping else ""
        if plain_text(raw_code).startswith(SUMMARY_STOP_PREFIXES):
            break
        if items and commercial_summary_boundary(row, mapping):
            break
        boundary = project_boundary_title(row, mapping)
        if boundary:
            pending_project_name = boundary
            continue
        description = safe_bc3_text(cell("description"))
        if not description:
            continue
        if len(header_map(row)) >= 2:
            continue
        # Las filas "Sub Total ..." son fórmulas de presentación, no conceptos
        # del itemizado ni capítulos editables.
        if plain_text(description).startswith(NON_CONCEPT_PREFIXES):
            continue
        sigla_match = re.search(r"\bSIGLA\s*[\"'“”]?\s*([A-Z]{1,3}\d?)\b", str(raw_code or ""), re.IGNORECASE)
        sigla_code = normalize_code(sigla_match.group(1)) if sigla_match else ""
        if sigla_code and sigla_code in regular_codes:
            continue
        code = sigla_code or normalize_code(raw_code)
        unit = safe_bc3_text(cell("unit")).upper()
        quantity = parse_number(cell("quantity"))
        price = parse_number(cell("price"))
        total = parse_number(cell("total"))
        if not code:
            # Sin código sólo puede crearse una partida si Unidad está informada.
            # Totales, IVA y fórmulas de resumen no son conceptos del itemizado.
            if not unit:
                continue
            code = str(auto_number)
            auto_number += 1
        if pending_project_name:
            if active_project_items:
                active_project_index += 1
                active_project_items = 0
            active_project_name = pending_project_name
            pending_project_name = ""
        items.append(Item(
            code, description, unit, quantity, price, total,
            source_row=row_number,
            project_index=active_project_index,
            project_name=active_project_name,
        ))
        active_project_items += 1
    return items, best_index, mapping


def assign_bc3_codes(items: Iterable[Item], warnings: list[str], reserved: Iterable[str] = ()) -> None:
    used: set[str] = {"ROOT##", *reserved}
    for item in items:
        candidate = item.code
        if len(candidate) > 12 or candidate in used or not candidate:
            digest = hashlib.sha1(item.code.encode("utf-8")).hexdigest().upper()[:8]
            candidate = f"IT{digest}"
            counter = 2
            while candidate in used:
                candidate = f"IT{digest[:6]}{counter:02d}"
                counter += 1
            warnings.append(f"Código '{item.code}' normalizado como '{candidate}' para compatibilidad BC3.")
        item.bc3_code = candidate
        used.add(candidate)


def normalize_bc3_ref(value: object) -> str:
    reference = re.sub(r"\s+", "", str(value or "")).upper()
    # ROOT## is CodeAPU's master budget sentinel; the last # is not merely the
    # FIEBDC chapter suffix and therefore must survive normalization.
    return "ROOT##" if reference == "ROOT##" else reference.rstrip("#")


def parse_bc3_number(value: object) -> float:
    """Parse FIEBDC-3 numeric fields without Excel thousands heuristics.

    BC3 writes decimal quantities with a dot (for example ``217.855``).  The
    generic spreadsheet parser intentionally interprets that shape as a
    thousands-grouped integer, which is correct for many Chilean Excel files
    but would multiply BC3 decompositions by 1,000.
    """
    text = str(value or "").strip().replace(" ", "")
    if not text:
        return 0.0
    text = re.sub(r"[^0-9,\.\-+]", "", text)
    if "," in text and "." in text:
        decimal = "," if text.rfind(",") > text.rfind(".") else "."
        thousands = "." if decimal == "," else ","
        text = text.replace(thousands, "").replace(decimal, ".")
    else:
        text = text.replace(",", ".")
    try:
        number = float(text)
        return number if math.isfinite(number) else 0.0
    except ValueError:
        return 0.0


def decode_bc3(data: bytes) -> tuple[str, str]:
    if data.startswith(b"\xef\xbb\xbf"):
        return data.decode("utf-8-sig"), "UTF-8"
    try:
        return data.decode("utf-8"), "UTF-8"
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace"), "Windows-1252"


def import_bc3_bytes(data: bytes, filename: str) -> ConversionResult:
    text, encoding = decode_bc3(data)
    concepts: dict[str, dict] = {}
    concept_order: list[str] = []
    relations: dict[str, list[tuple[str, float, float]]] = {}
    warnings: list[str] = []
    version = "FIEBDC-3"

    for raw_line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = raw_line.strip("\ufeff\x00 ")
        if not line.startswith("~"):
            continue
        fields = line.split("|")
        record = fields[0].upper()
        if record == "~V":
            version = safe_bc3_text(fields[2] if len(fields) > 2 else fields[0]) or version
        elif record == "~C" and len(fields) >= 4:
            raw_ref = fields[1].strip()
            code = normalize_bc3_ref(raw_ref)
            if not code:
                continue
            if code not in concepts:
                concept_order.append(code)
            concepts[code] = {
                "code": code,
                "chapterRef": raw_ref.rstrip().endswith("#"),
                "unit": safe_bc3_text(fields[2]).upper(),
                "description": safe_bc3_text(fields[3]) or code,
                "price": parse_bc3_number(fields[4] if len(fields) > 4 else 0),
                "type": safe_bc3_text(fields[6] if len(fields) > 6 else ""),
            }
        elif record == "~D" and len(fields) >= 3:
            parent = normalize_bc3_ref(fields[1])
            tokens = [token for token in fields[2].split("\\") if token != ""]
            rows = relations.setdefault(parent, [])
            for index in range(0, len(tokens) - 2, 3):
                child = normalize_bc3_ref(tokens[index])
                if child:
                    factor = parse_bc3_number(tokens[index + 1])
                    yield_value = parse_bc3_number(tokens[index + 2])
                    # FIEBDC-3 ~D stores factor and yield as separate fields;
                    # the effective quantity is their product. Preserve both
                    # values so CodeAPU can edit and re-export them independently.
                    stored_factor = factor if factor != 0 else 1.0
                    stored_yield = yield_value if yield_value != 0 else 1.0
                    rows.append((child, stored_factor, stored_yield))

    if not concepts:
        raise ValueError("El archivo BC3 no contiene registros ~C reconocibles.")
    missing = sorted({child for rows in relations.values() for child, _, _ in rows if child not in concepts})
    if missing:
        warnings.append(f"Se omitieron {len(missing)} referencias sin concepto: {', '.join(missing[:5])}.")
        for parent in relations:
            relations[parent] = [
                (child, factor, yield_value)
                for child, factor, yield_value in relations[parent]
                if child in concepts
            ]

    child_codes = {child for rows in relations.values() for child, _, _ in rows}
    master_code = "ROOT##" if "ROOT##" in concepts else ""
    project_specs: list[tuple[str, str, list[tuple[str, float, float]], bool]] = []

    if master_code:
        root_children = relations.get(master_code, [])
        explicit_projects = [entry for entry in root_children if re.fullmatch(r"PRJ\d+", entry[0])]
        if explicit_projects and len(explicit_projects) == len(root_children):
            for code, _, _ in explicit_projects:
                project_specs.append((code, concepts[code]["description"], relations.get(code, []), True))
        else:
            project_specs.append(("PRJ001", concepts[master_code]["description"], root_children, False))
    else:
        graph_roots = [code for code in concept_order if code not in child_codes and (relations.get(code) or not concepts[code]["unit"])]
        if not graph_roots:
            graph_roots = [concept_order[0]]
        for index, code in enumerate(graph_roots, start=1):
            project_specs.append((f"PRJ{index:03d}", concepts[code]["description"], relations.get(code, []), True))

    project_roots: list[dict] = []
    items: list[dict] = []
    analyses: dict[str, list[dict]] = {}
    budget_codes: set[str] = set()
    analysis_resource_codes: set[str] = set()
    commercial_labels = (
        "costo total", "gastos generales", "utilidad", "valor neto",
        "iva", "i v a", "valor total presupuesto", "total oferta",
    )

    def build_analysis(code: str, stack: tuple[str, ...] = ()) -> None:
        if code in stack or code in analyses or not relations.get(code):
            return
        rows: list[dict] = []
        for child, factor, quantity in relations.get(code, []):
            concept = concepts.get(child)
            if not concept:
                continue
            is_subanalysis = bool(relations.get(child))
            rows.append({
                "code": child,
                "description": concept["description"],
                "unit": concept["unit"] or ("UN" if is_subanalysis else ""),
                "quantity": quantity,
                "factor": factor,
                "unitPrice": concept["price"],
                "isSubanalysis": is_subanalysis,
            })
            analysis_resource_codes.add(child)
            if is_subanalysis:
                build_analysis(child, (*stack, code))
        if rows:
            analyses[code] = rows

    def walk_budget(
        entries: list[tuple[str, float, float]], project_code: str, project_name: str,
        parent_code: str, level: int, top_code: str = "", path: tuple[str, ...] = (),
    ) -> None:
        for code, factor, yield_value in entries:
            if code in path:
                warnings.append(f"Se evitó una relación circular en '{code}'.")
                continue
            concept = concepts.get(code)
            if not concept:
                continue
            quantity = factor * yield_value
            is_partida = bool(concept["unit"])
            current_top = top_code or code
            top_description = plain_text(concepts.get(current_top, concept)["description"])
            root_kind = "commercial" if any(label in top_description for label in commercial_labels) else "budget"
            item = {
                "code": code,
                "originalCode": code,
                "description": concept["description"],
                "unit": concept["unit"],
                "quantity": quantity,
                "price": concept["price"],
                "total": quantity * concept["price"],
                "level": level,
                "synthetic": False,
                "hasChildren": bool(relations.get(code)) and not is_partida,
                "isChapter": not is_partida,
                "isPartida": is_partida,
                "parentCode": parent_code,
                "rootCode": current_top,
                "rootKind": root_kind,
                "projectCode": project_code,
                "projectName": project_name,
            }
            items.append(item)
            budget_codes.add(code)
            if is_partida:
                build_analysis(code)
            else:
                walk_budget(relations.get(code, []), project_code, project_name, code, level + 1, current_top, (*path, code))

    for project_code, project_name, root_entries, existing_concept in project_specs:
        project_price = concepts.get(project_code, {}).get("price", 0.0) if existing_concept else 0.0
        if not project_price:
            project_price = sum(
                concepts.get(child, {}).get("price", 0.0) * factor * yield_value
                for child, factor, yield_value in root_entries
            )
        project_roots.append({"code": project_code, "description": project_name, "total": project_price, "level": 1})
        walk_budget(root_entries, project_code, project_name, project_code, 2)

    if not items:
        raise ValueError("El BC3 no contiene una jerarquía de presupuesto utilizable.")
    root_total = concepts.get(master_code, {}).get("price", 0.0) if master_code else sum(root["total"] for root in project_roots)
    project_name = concepts.get(master_code, {}).get("description", "") or (project_roots[0]["description"] if len(project_roots) == 1 else Path(filename).stem)
    leaf_count = sum(1 for item in items if item["isPartida"])
    chapter_count = len(items) - leaf_count
    import_payload = {
        "sourceFormat": "BC3",
        "projectRoot": project_roots[0],
        "projectRoots": project_roots,
        "projectCount": len(project_roots),
        "items": items,
        "analyses": analyses,
        "bc3ConceptCount": len(concepts),
    }
    return ConversionResult(
        source_name=filename,
        sheet_name=f"{version} · {encoding}",
        project_name=project_name,
        items=[],
        warnings=warnings,
        bc3_bytes=data,
        root_total=root_total,
        leaf_count=leaf_count,
        chapter_count=chapter_count,
        project_roots=project_roots,
        import_payload=import_payload,
        normalization_report={
            "selectionMethod": "registros FIEBDC-3 ~C y ~D",
            "candidateSheets": 0,
            "headerRow": "BC3",
            "detectedColumns": ["~C conceptos", "~D descomposiciones"],
            "sourceConcepts": len(concepts),
            "outputConcepts": len(items),
            "syntheticParents": 0,
            "duplicateCodesPreserved": 0,
            "bc3CodesRemapped": 0,
            "orderPreserved": True,
            "classificationRule": "Unidad vacía = capítulo; Unidad informada = partida/APU",
            "projectCount": len(project_roots),
            "analysisCount": sum(1 for rows in analyses.values() if rows),
            "resourceCount": len(analysis_resource_codes),
            "encoding": encoding,
        },
    )


def convert_bytes(data: bytes, filename: str, project_name: str = "") -> ConversionResult:
    suffix = Path(filename).suffix.lower()
    if suffix == ".bc3":
        return import_bc3_bytes(data, filename)
    sheets = read_workbook(data, suffix)
    candidates: list[tuple[float, int, str, list[Item], int, dict[str, int]]] = []
    for sheet_name, rows in sheets:
        parsed, header_index, mapping = parse_sheet(rows)
        unit_rows = sum(1 for item in parsed if item.unit.strip())
        hierarchical_rows = sum(1 for item in parsed if "." in item.code)
        # Prima una tabla utilizable (Unidad + jerarquía + columnas reconocidas)
        # por sobre una hoja resumen que sólo tenga muchas descripciones.
        quality = len(parsed) + unit_rows * 3 + hierarchical_rows * 0.25 + len(mapping) * 25
        candidates.append((quality, len(parsed), sheet_name, parsed, header_index, mapping))
    _, count, sheet_name, parsed_items, header_index, selected_mapping = max(
        candidates, key=lambda entry: entry[0], default=(0, 0, "", [], -1, {})
    )
    if not count:
        raise ValueError("No se encontró una tabla con descripción de partidas y datos reconocibles.")

    warnings: list[str] = []
    by_code: dict[str, Item] = {}
    source_order: list[str] = []
    root_occurrences: dict[tuple[int, str], int] = {}
    active_root_occurrence: dict[tuple[int, str], int] = {}
    duplicate_count = 0
    root_projects: dict[str, set[int]] = {}
    for parsed in parsed_items:
        parsed_root = parsed.code.split(".")[0]
        root_projects.setdefault(parsed_root, set()).add(parsed.project_index)
    for item in parsed_items:
        original = item.code
        parts = original.split(".")
        root = parts[0]
        root_key = (item.project_index, root)
        if len(parts) == 1:
            root_occurrences[root_key] = root_occurrences.get(root_key, 0) + 1
            active_root_occurrence[root_key] = root_occurrences[root_key]
        occurrence = active_root_occurrence.get(root_key, 1)
        first_project_for_root = min(root_projects.get(root, {item.project_index}))
        namespaced_root = root if item.project_index == first_project_for_root else f"P{item.project_index}{root}"
        internal_root = namespaced_root if occurrence == 1 else f"{namespaced_root}_R{occurrence}"
        internal = internal_root if len(parts) == 1 else ".".join([internal_root, *parts[1:]])
        was_duplicate = occurrence > 1
        if internal in by_code:
            duplicate = 2
            base = internal
            while f"{base}_D{duplicate}" in by_code:
                duplicate += 1
            internal = f"{base}_D{duplicate}"
            was_duplicate = True
        item.original_code = original
        item.code = internal
        if was_duplicate:
            duplicate_count += 1
            warnings.append(f"Código repetido '{original}' conservado como concepto único '{internal}'.")
        by_code[internal] = item
        source_order.append(internal)

    available = set(by_code)
    for code in list(available):
        parts = code.split(".")
        while len(parts) > 1:
            parts.pop()
            parent = ".".join(parts)
            if parent not in by_code:
                descendant = by_code[code]
                original_parent = re.sub(r"_(?:P|R)\d+(?=\.|$)", "", parent)
                if descendant.project_index > 1:
                    original_parent = re.sub(rf"^P{descendant.project_index}(?=[A-Z0-9])", "", original_parent)
                by_code[parent] = Item(
                    parent, f"Capítulo {original_parent}", synthetic=True, original_code=original_parent,
                    source_row=descendant.source_row,
                    project_index=descendant.project_index,
                    project_name=descendant.project_name,
                )
                available.add(parent)
                warnings.append(f"Se creó el capítulo faltante '{parent}'.")

    # Conserva el orden entregado por el mandante. Los capítulos sintéticos se
    # insertan justo antes de su primer descendiente, no mediante orden alfabético.
    codes: list[str] = []
    seen: set[str] = set()
    for source_code in source_order:
        parts = source_code.split(".")
        chain = [".".join(parts[:index]) for index in range(1, len(parts) + 1)]
        for code in chain:
            if code in by_code and code not in seen:
                codes.append(code)
                seen.add(code)
    for code in by_code:
        if code not in seen:
            codes.append(code)
            seen.add(code)
    title = safe_bc3_text(project_name) or safe_bc3_text(Path(filename).stem)
    project_indexes = list(dict.fromkeys(by_code[code].project_index for code in codes))
    project_codes = {index: f"PRJ{index:03d}" for index in project_indexes}
    project_names: dict[int, str] = {}
    for index in project_indexes:
        detected = next((by_code[code].project_name for code in codes if by_code[code].project_index == index and by_code[code].project_name), "")
        project_names[index] = safe_bc3_text(detected) or (title if len(project_indexes) == 1 else f"{title} · Proyecto {index}")

    children: dict[str, list[str]] = {}
    for code in codes:
        parent = parent_code(code, available)
        project_parent = project_codes[by_code[code].project_index]
        children.setdefault(parent or project_parent, []).append(code)

    prices: dict[str, float] = {}
    quantities: dict[str, float] = {}
    for code in reversed(codes):
        item = by_code[code]
        if children.get(code):
            quantities[code] = 1.0
            prices[code] = sum(prices[child] * quantities[child] for child in children[code])
            item.price = prices[code]
            item.quantity = 1.0
            item.total = prices[code]
        else:
            quantity = item.quantity or (1.0 if item.price or item.total else 0.0)
            price = item.price or (item.total / quantity if item.total and quantity else 0.0)
            quantities[code], prices[code] = quantity, price
            item.quantity, item.price = quantity, price
            item.total = item.total or quantity * price

    ordered_items = [by_code[code] for code in codes]
    assign_bc3_codes(ordered_items, warnings, project_codes.values())
    code_map = {item.code: item.bc3_code for item in ordered_items}

    def concept_code(code: str) -> str:
        base = code_map[code]
        return f"{base}#" if children.get(code) else base

    project_totals = {
        index: sum(prices[child] * quantities[child] for child in children.get(project_codes[index], []))
        for index in project_indexes
    }
    project_roots = [
        {"code": project_codes[index], "description": project_names[index], "total": project_totals[index]}
        for index in project_indexes
    ]
    root_total = sum(project_totals.values())
    today = datetime.now().strftime("%d%m%y")
    lines = [
        "~V|CodeAPU CODEAPU.CL|FIEBDC-3/2002|Gestor de APU compatible con Presto 8.8||ANSI|",
        "~K|\\2\\2\\3\\0\\0\\0\\0\\CLP\\|0|",
        f"~C|ROOT##||{title}|{format_number(root_total)}|{today}|0|",
    ]
    for root in project_roots:
        lines.append(f"~C|{root['code']}#||{root['description']}|{format_number(root['total'])}|{today}|0|")
    for item in ordered_items:
        unit = "" if children.get(item.code) else item.unit
        lines.append(
            f"~C|{concept_code(item.code)}|{unit}|{item.description}|{format_number(prices[item.code])}|{today}|0|"
        )
    root_tokens: list[str] = []
    for root in project_roots:
        root_tokens.extend([f"{root['code']}#", "1", "1"])
    lines.append(f"~D|ROOT##|{'\\'.join(root_tokens)}|")
    for parent in [*project_codes.values(), *codes]:
        if not children.get(parent):
            continue
        tokens: list[str] = []
        for child in children[parent]:
            tokens.extend([concept_code(child), "1", format_number(quantities[child])])
        parent_ref = f"{parent}#" if parent in project_codes.values() else concept_code(parent)
        lines.append(f"~D|{parent_ref}|{'\\'.join(tokens)}|")

    bc3_text = "\r\n".join(lines) + "\r\n"
    bc3_bytes = bc3_text.encode("cp1252", errors="replace")
    result = ConversionResult(
        source_name=filename,
        sheet_name=sheet_name,
        project_name=title,
        items=ordered_items,
        warnings=warnings,
        bc3_bytes=bc3_bytes,
        root_total=root_total,
        # Regla funcional CodeAPU: la columna Unidad determina la naturaleza.
        leaf_count=sum(1 for item in ordered_items if item.unit.strip()),
        chapter_count=sum(1 for item in ordered_items if not item.unit.strip()),
        normalization_report={
            "selectionMethod": "calidad estructural",
            "candidateSheets": len(candidates),
            "headerRow": header_index + 1,
            "detectedColumns": sorted(selected_mapping),
            "sourceConcepts": count,
            "outputConcepts": len(ordered_items),
            "syntheticParents": sum(1 for item in ordered_items if item.synthetic),
            "duplicateCodesPreserved": duplicate_count,
            "bc3CodesRemapped": sum(1 for item in ordered_items if item.bc3_code != item.code),
            "orderPreserved": True,
            "classificationRule": "Unidad vacía = capítulo; Unidad informada = partida",
            "projectCount": len(project_roots),
            "projectDetectionRule": "Rótulos de obra/itemizado/edificio/sede y reinicio de bloques jerárquicos",
        },
        project_roots=project_roots,
    )
    validate_result(result)
    return result


def validate_result(result: ConversionResult) -> None:
    text = result.bc3_bytes.decode("cp1252")
    concepts: set[str] = set()
    relations: list[tuple[str, str]] = []
    for line in text.splitlines():
        if line.startswith("~C|"):
            concepts.add(line.split("|")[1])
        elif line.startswith("~D|"):
            parts = line.split("|")
            tokens = [token for token in parts[2].split("\\") if token]
            relations.extend((parts[1], tokens[index]) for index in range(0, len(tokens), 3))
    missing = sorted({ref for pair in relations for ref in pair if ref not in concepts})
    if missing:
        raise ValueError(f"Validación BC3 fallida; referencias sin concepto: {', '.join(missing[:5])}")
    if any(len(code.rstrip("#")) > 13 for code in concepts):
        raise ValueError("Validación BC3 fallida; existe un código de más de 13 caracteres.")

